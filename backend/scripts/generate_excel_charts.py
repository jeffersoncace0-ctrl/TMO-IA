#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador nativo de Excel con 4 hojas y gráficos integrados (BarChart y DoughnutChart)
utilizando openpyxl para Plataforma de Análisis de Llamadas y TMO.
"""

import sys
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, DoughnutChart, Reference, Series
from openpyxl.utils import get_column_letter

def create_styled_cell(cell, value, font_color="000000", bg_color=None, bold=False, align="left", size=11, italic=False):
    cell.value = value
    cell.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=font_color)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if bg_color:
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

def apply_borders(ws, start_row, end_row, start_col, end_col):
    thin = Side(border_style="thin", color="D3D3D3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(row=r, column=c).border = border

def generate_report(data_json_path, output_excel_path):
    with open(data_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    wb = Workbook()
    
    # ---------------------------------------------------------
    # HOJA 1: Datos Originales
    # ---------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Datos Originales"
    ws1.views.sheetView[0].showGridLines = True

    # Encabezado informativo
    ws1.merge_cells("A1:G1")
    create_styled_cell(ws1["A1"], "REGISTRO DE DATOS CRUDOS IMPORTADOS - CALL CENTER", font_color="FFFFFF", bg_color="1E3A8A", bold=True, size=13)
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    create_styled_cell(ws1["A2"], f"Archivo: {data.get('fileName', 'N/A')} | Fecha Proceso: {data.get('processedAt', 'N/A')} | Filas: {len(data.get('calls', []))}", italic=True, size=10)
    ws1.merge_cells("A2:G2")

    headers1 = ["ID Llamada", "Código Agente", "Nombre Agente", "Fecha", "Tipo", "Estado", "Duración (s)", "Conversación (s)", "Silencio (s)", "Válido"]
    for col_idx, h in enumerate(headers1, 1):
        create_styled_cell(ws1.cell(row=4, column=col_idx), h, font_color="FFFFFF", bg_color="3B82F6", bold=True, align="center")
    ws1.row_dimensions[4].height = 24

    for row_idx, call in enumerate(data.get("calls", []), 5):
        create_styled_cell(ws1.cell(row=row_idx, column=1), call.get("originalCallId", f"CALL-{row_idx-4}"), align="center")
        create_styled_cell(ws1.cell(row=row_idx, column=2), call.get("agentCode", ""), align="center")
        create_styled_cell(ws1.cell(row=row_idx, column=3), call.get("agentName", ""))
        create_styled_cell(ws1.cell(row=row_idx, column=4), str(call.get("callDate", "")).replace("T", " ")[:19], align="center")
        create_styled_cell(ws1.cell(row=row_idx, column=5), call.get("callType", "ENTRANTE"), align="center")
        create_styled_cell(ws1.cell(row=row_idx, column=6), call.get("callStatus", "CONTESTADA"), align="center")
        create_styled_cell(ws1.cell(row=row_idx, column=7), round(call.get("totalDurationSeconds", 0), 2), align="right")
        create_styled_cell(ws1.cell(row=row_idx, column=8), round(call.get("talkTimeSeconds", 0), 2), align="right")
        create_styled_cell(ws1.cell(row=row_idx, column=9), round(call.get("silenceTimeSeconds", 0), 2), align="right")
        
        is_val = call.get("isValid", True)
        val_text = "SÍ" if is_val else "NO"
        val_bg = "DCFCE7" if is_val else "FEE2E2"
        val_color = "166534" if is_val else "991B1B"
        create_styled_cell(ws1.cell(row=row_idx, column=10), val_text, font_color=val_color, bg_color=val_bg, align="center", bold=True)
        ws1.row_dimensions[row_idx].height = 20

    apply_borders(ws1, 4, 4 + len(data.get("calls", [])), 1, 10)

    # ---------------------------------------------------------
    # HOJA 2: Resumen de Agentes
    # ---------------------------------------------------------
    ws2 = wb.create_sheet(title="Resumen de Agentes")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:H1")
    create_styled_cell(ws2["A1"], f"CONSOLIDADO DE AGENTES Y TMO (UMBRAL: {data.get('thresholdSeconds', 20)}s)", font_color="FFFFFF", bg_color="1E3A8A", bold=True, size=13)
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    headers2 = ["Código", "Nombre del Agente", "Total Llamadas", "Gestiones Válidas", "TMO Promedio (s)", "Silencio Promedio (s)", "% Cumplimiento", "Clasificación"]
    for col_idx, h in enumerate(headers2, 1):
        create_styled_cell(ws2.cell(row=3, column=col_idx), h, font_color="FFFFFF", bg_color="3B82F6", bold=True, align="center")
    ws2.row_dimensions[3].height = 24

    agents = data.get("agentSummaries", [])
    for idx, ag in enumerate(agents, 4):
        create_styled_cell(ws2.cell(row=idx, column=1), ag.get("agentCode", ""), align="center")
        create_styled_cell(ws2.cell(row=idx, column=2), ag.get("agentName", ""))
        create_styled_cell(ws2.cell(row=idx, column=3), ag.get("totalCalls", 0), align="right")
        create_styled_cell(ws2.cell(row=idx, column=4), ag.get("validCalls", 0), align="right")
        create_styled_cell(ws2.cell(row=idx, column=5), round(ag.get("ahtTmoSeconds", 0), 2), align="right", bold=True)
        create_styled_cell(ws2.cell(row=idx, column=6), round(ag.get("avgSilenceTimeSeconds", 0), 2), align="right")
        create_styled_cell(ws2.cell(row=idx, column=7), f"{round(ag.get('complianceRate', 0), 1)}%", align="right")

        clasif = ag.get("classification", "REQUIERE_REVISION")
        bg_col = "DCFCE7" if clasif == "APROBADO" else ("FEE2E2" if clasif == "REQUIERE_REVISION" else "F3F4F6")
        fg_col = "166534" if clasif == "APROBADO" else ("991B1B" if clasif == "REQUIERE_REVISION" else "4B5563")
        create_styled_cell(ws2.cell(row=idx, column=8), clasif, font_color=fg_col, bg_color=bg_col, align="center", bold=True)
        ws2.row_dimensions[idx].height = 20

    apply_borders(ws2, 3, 3 + len(agents), 1, 8)

    # GRÁFICO DE BARRAS NATIVO EN HOJA 2: TMO por Agente
    if len(agents) > 0:
        chart1 = BarChart()
        chart1.type = "col"
        chart1.style = 10
        chart1.title = "TMO Promedio por Agente (s)"
        chart1.y_axis.title = "Segundos"
        chart1.x_axis.title = "Agente"
        chart1.width = 16
        chart1.height = 10

        data_ref = Reference(ws2, min_col=5, min_row=3, max_row=3 + len(agents))
        cats_ref = Reference(ws2, min_col=1, min_row=4, max_row=3 + len(agents))
        chart1.add_data(data_ref, titles_from_data=True)
        chart1.set_categories(cats_ref)
        chart1.legend = None
        ws2.add_chart(chart1, "J3")

    # ---------------------------------------------------------
    # HOJA 3: Análisis Estadístico
    # ---------------------------------------------------------
    ws3 = wb.create_sheet(title="Análisis Estadístico")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:E1")
    create_styled_cell(ws3["A1"], "ANÁLISIS ESTADÍSTICO Y DISTRIBUCIONES OPERATIVAS", font_color="FFFFFF", bg_color="1E3A8A", bold=True, size=13)
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 28

    kpis = data.get("kpis", {})
    kpi_rows = [
        ("Total de Llamadas", kpis.get("totalCalls", 0)),
        ("Total de Agentes Únicos", kpis.get("totalAgents", 0)),
        ("TMO Global Ponderado (s)", round(kpis.get("averageHandleTimeSeconds", 0), 2)),
        ("Tiempo Promedio Conversación (s)", round(kpis.get("averageTalkTimeSeconds", 0), 2)),
        ("Tiempo Promedio Silencio (s)", round(kpis.get("averageSilenceTimeSeconds", 0), 2)),
        ("Porcentaje de Cumplimiento Global", f"{round(kpis.get('percentageOfCompliance', 0), 1)}%"),
        ("Umbral Aplicado (s)", f"{kpis.get('thresholdApplied', 20)}s"),
        ("Agentes Aprobados", kpis.get("approvedAgentsCount", 0)),
        ("Agentes Requieren Revisión", kpis.get("reviewRequiredAgentsCount", 0)),
        ("Agentes con Error de Datos", kpis.get("errorAgentsCount", 0))
    ]

    create_styled_cell(ws3["A3"], "Métrica Global", font_color="FFFFFF", bg_color="3B82F6", bold=True)
    create_styled_cell(ws3["B3"], "Valor", font_color="FFFFFF", bg_color="3B82F6", bold=True, align="right")
    ws3.row_dimensions[3].height = 22

    for i, (label, val) in enumerate(kpi_rows, 4):
        create_styled_cell(ws3.cell(row=i, column=1), label)
        create_styled_cell(ws3.cell(row=i, column=2), val, align="right", bold=True)
        ws3.row_dimensions[i].height = 20
    apply_borders(ws3, 3, 3 + len(kpi_rows), 1, 2)

    # Tabla de Distribución de Estado de Llamadas
    create_styled_cell(ws3["D3"], "Estado Llamada", font_color="FFFFFF", bg_color="3B82F6", bold=True)
    create_styled_cell(ws3["E3"], "Cantidad", font_color="FFFFFF", bg_color="3B82F6", bold=True, align="right")
    
    create_styled_cell(ws3["D4"], "Contestadas")
    create_styled_cell(ws3["E4"], kpis.get("acceptedCalls", 0), align="right")
    create_styled_cell(ws3["D5"], "Perdidas")
    create_styled_cell(ws3["E5"], kpis.get("missedCalls", 0), align="right")
    apply_borders(ws3, 3, 5, 4, 5)

    # GRÁFICO DE DONA NATIVO EN HOJA 3: Estado de Llamadas
    chart2 = DoughnutChart()
    chart2.title = "Distribución: Contestadas vs Perdidas"
    chart2.style = 10
    chart2.width = 14
    chart2.height = 9

    labels_ref = Reference(ws3, min_col=4, min_row=4, max_row=5)
    data_ref2 = Reference(ws3, min_col=5, min_row=3, max_row=5)
    chart2.add_data(data_ref2, titles_from_data=True)
    chart2.set_categories(labels_ref)
    ws3.add_chart(chart2, "G3")

    # ---------------------------------------------------------
    # HOJA 4: Validación de IA
    # ---------------------------------------------------------
    ws4 = wb.create_sheet(title="Validación de IA")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:G1")
    create_styled_cell(ws4["A1"], "AUDITORÍA INTELIGENTE Y EXPLICABILIDAD DE RESULTADOS (IA)", font_color="FFFFFF", bg_color="1E3A8A", bold=True, size=13)
    ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 28

    headers4 = ["Agente", "TMO (s)", "Umbral", "Regla Determinista", "Revisión IA", "Motivo de Clasificación", "Recomendación Operativa"]
    for col_idx, h in enumerate(headers4, 1):
        create_styled_cell(ws4.cell(row=3, column=col_idx), h, font_color="FFFFFF", bg_color="3B82F6", bold=True, align="center")
    ws4.row_dimensions[3].height = 24

    ai_items = data.get("aiValidations", [])
    for idx, item in enumerate(ai_items, 4):
        create_styled_cell(ws4.cell(row=idx, column=1), f"{item.get('agentName', '')} ({item.get('agentCode', '')})")
        create_styled_cell(ws4.cell(row=idx, column=2), round(item.get("ahtTmoSeconds", 0), 2), align="right")
        create_styled_cell(ws4.cell(row=idx, column=3), f"{item.get('thresholdApplied', 20)}s", align="center")
        
        regla = item.get("ruleResult", "REQUIERE_REVISION")
        regla_bg = "DCFCE7" if regla == "APROBADO" else "FEE2E2"
        regla_fg = "166534" if regla == "APROBADO" else "991B1B"
        create_styled_cell(ws4.cell(row=idx, column=4), regla, font_color=regla_fg, bg_color=regla_bg, align="center", bold=True)
        
        create_styled_cell(ws4.cell(row=idx, column=5), item.get("aiClassification", "CONFIRMADO"), align="center")
        create_styled_cell(ws4.cell(row=idx, column=6), item.get("classificationReason", ""))
        create_styled_cell(ws4.cell(row=idx, column=7), item.get("operationalRecommendation", ""))
        ws4.row_dimensions[idx].height = 36

    apply_borders(ws4, 3, 3 + len(ai_items), 1, 7)

    # Ajuste automático del ancho de columnas para todas las hojas
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Evitar que las celdas combinadas del título distorsionen el ancho
                if cell.row in [1, 2]:
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Ajustes específicos para columnas de texto largo en Hoja 4
    ws4.column_dimensions["F"].width = 45
    ws4.column_dimensions["G"].width = 45

    wb.save(output_excel_path)
    print(f"SUCCESS: Reporte generado exitosamente con gráficos nativos en: {output_excel_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python generate_excel_charts.py <input_data.json> <output_report.xlsx>")
        sys.exit(1)
    generate_report(sys.argv[1], sys.argv[2])
